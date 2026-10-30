"""Audit the electricity technology split in the skill's bundled CHEER-Clean workbook.

This is a benchmark accounting check, not a CGE solve.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[5]
DEFAULT_WORKBOOK = (Path(__file__).resolve().parents[1] / "assets" /
                    "cheer_clean_electricity_example" / "v1" / "benchmark.xlsx")
TOLERANCE = 1e-4  # workbook's native monetary units
GROUPS = (("AgriPr", "FrsPr"), ("Biofuel", "OthMfg"))


def value(cell: object, label: str) -> float:
    if cell is None:
        return 0.0
    if isinstance(cell, bool) or not isinstance(cell, (int, float)):
        raise ValueError(f"{label}: expected a numeric value or blank cell")
    result = float(cell)
    if not math.isfinite(result):
        raise ValueError(f"{label}: nonfinite value")
    return result


def labels_in_column(sheet: object, column: int, end_label: str) -> dict[str, int]:
    labels: dict[str, int] = {}
    found_end = False
    for row in range(2, sheet.max_row + 1):
        raw = sheet.cell(row, column).value
        if raw is None:
            continue
        label = str(raw).strip()
        if label in labels:
            raise ValueError(f"{sheet.title}: duplicate row label {label}")
        labels[label] = row
        if label == end_label:
            found_end = True
            break
    if not found_end:
        raise ValueError(f"{sheet.title}: missing {end_label} row")
    return labels


def analyze(workbook: Path) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, object]]:
    if not workbook.is_file() or workbook.suffix.lower() != ".xlsx":
        raise ValueError("the bundled workbook must be an existing .xlsx file")
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        if not {"IO", "ELEC"}.issubset(book.sheetnames):
            raise ValueError("workbook must contain IO and ELEC sheets")
        io, elec = book["IO"], book["ELEC"]
        io_header = {str(io.cell(1, col).value).strip(): col for col in range(2, io.max_column + 1)}
        if "Elec" not in io_header or "SUM_OUTPUT" not in io_header:
            raise ValueError("IO: expected Elec and SUM_OUTPUT columns")
        parent_column = io_header["Elec"]
        io_rows = labels_in_column(io, 1, "tax")
        elec_rows = labels_in_column(elec, 1, "tax")
        if set(io_rows) != set(elec_rows) or not {"Elec", "labor", "capital", "tax"}.issubset(io_rows):
            raise ValueError("IO and ELEC benchmark input row labels differ")

        technology_columns: dict[str, int] = {}
        for col in range(2, elec.max_column + 1):
            raw = elec.cell(1, col).value
            if raw is None:
                break  # later columns are workbook notes, not technology accounts
            name = str(raw).strip()
            if name in technology_columns:
                raise ValueError(f"ELEC: duplicate technology {name}")
            technology_columns[name] = col
        if len(technology_columns) < 2:
            raise ValueError("ELEC: expected at least two technology columns")

        parent_by_row = {label: value(io.cell(row, parent_column).value, f"IO/{label}")
                         for label, row in io_rows.items()}
        child_by_row = {label: sum(value(elec.cell(elec_rows[label], col).value,
                                          f"ELEC/{label}/{tech}")
                                   for tech, col in technology_columns.items())
                        for label in io_rows}
        parent_total = sum(parent_by_row.values())
        child_total = sum(child_by_row.values())
        if parent_total <= 0 or abs(child_total - parent_total) > TOLERANCE:
            raise ValueError("technology input total does not reproduce the IO electricity column")

        technologies: list[dict[str, object]] = []
        for tech, col in technology_columns.items():
            labor = value(elec.cell(elec_rows["labor"], col).value, f"ELEC/labor/{tech}")
            capital = value(elec.cell(elec_rows["capital"], col).value, f"ELEC/capital/{tech}")
            tax = value(elec.cell(elec_rows["tax"], col).value, f"ELEC/tax/{tech}")
            intermediate = sum(value(elec.cell(row, col).value, f"ELEC/{label}/{tech}")
                               for label, row in elec_rows.items()
                               if label not in {"labor", "capital", "tax"})
            total = intermediate + labor + capital + tax
            if total <= 0 or labor < 0 or capital < 0 or intermediate < 0:
                raise ValueError(f"ELEC/{tech}: invalid benchmark cost components")
            technologies.append({"technology": tech, "intermediate": intermediate,
                                 "labor": labor, "capital": capital, "net_tax": tax,
                                 "benchmark_value": total,
                                 "share_of_electricity_value": total / parent_total})

        rows = [{"input": label, "io_electricity": parent_by_row[label],
                 "sum_technologies": child_by_row[label],
                 "difference": child_by_row[label] - parent_by_row[label]}
                for label in io_rows]
        grouped = set()
        group_residuals: dict[str, float] = {}
        for group in GROUPS:
            if set(group).issubset(io_rows):
                group_residuals["+".join(group)] = sum(child_by_row[k] - parent_by_row[k] for k in group)
                grouped.update(group)
        unexplained = {label: child_by_row[label] - parent_by_row[label]
                       for label in io_rows if label not in grouped
                       and abs(child_by_row[label] - parent_by_row[label]) > TOLERANCE}
        if unexplained or any(abs(x) > TOLERANCE for x in group_residuals.values()):
            raise ValueError("unexplained input-row differences between IO and ELEC")

        sum_input_rows = labels_in_column(io, 1, "SUM_INPUT")
        reported_input = value(io.cell(sum_input_rows["SUM_INPUT"], parent_column).value,
                               "IO/SUM_INPUT/Elec")
        reported_output = value(io.cell(io_rows["Elec"], io_header["SUM_OUTPUT"]).value,
                                "IO/Elec/SUM_OUTPUT")
        if any(abs(x - parent_total) > TOLERANCE for x in (reported_input, reported_output)):
            raise ValueError("IO electricity totals are not balanced")
        checksum = hashlib.sha256(workbook.read_bytes()).hexdigest()
        validation = {"analysis_level": "benchmark_accounting_only_no_cge_solve",
                      "workbook_name": workbook.name, "sha256": checksum,
                      "technology_count": len(technologies),
                      "parent_io_value": parent_total, "technology_value_total": child_total,
                      "total_difference": child_total - parent_total,
                      "grouped_row_differences": group_residuals,
                      "unexplained_row_differences": unexplained,
                      "workbook_units": "native input workbook units; no unit conversion in this audit"}
        return technologies, rows, validation
    finally:
        book.close()


def write_csv(path: Path, records: list[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out-dir", required=True, type=Path, help="new result directory outside Git")
    args = parser.parse_args()
    out_dir = args.out_dir.resolve()
    if out_dir.is_relative_to(REPO_ROOT) or any((p / ".git").exists() for p in (out_dir, *out_dir.parents)):
        parser.error("--out-dir must be outside all Git worktrees")
    if out_dir.exists() or not out_dir.parent.is_dir():
        parser.error("--out-dir must be new and its parent must exist")
    technologies, rows, validation = analyze(DEFAULT_WORKBOOK)
    out_dir.mkdir()
    write_csv(out_dir / "electricity_technologies.csv", technologies)
    write_csv(out_dir / "input_reconciliation.csv", rows)
    (out_dir / "validation.json").write_text(json.dumps(validation, indent=2, ensure_ascii=False) + "\n",
                                              encoding="utf-8")
    print(json.dumps({"result_directory": str(out_dir), **validation}, ensure_ascii=False))


if __name__ == "__main__":
    main()
